#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""配表正文的取数：把工作簿字节读成工作表文本，并按字符额度精确收敛。

## 为什么单独一个文件

`services/ai/platform_provider.py` 贴着仓库的长度闸门（`scripts/check_file_length.py --strict`
在 2000 行报错）。而「打开工作簿、按工作表切、把统计与正文装配进一个字符额度」这一块与
「向模型交付什么」的其余部分几乎没有耦合：输入是文件字节 + 仓库的两个表头配置，输出是
一段文本 —— 不碰数据库、不碰 Flask，也不认识 `ContextProvider` 那套契约。搬出来之后它能
被完整单测，`platform_provider` 只保留一行回导（旧名字仍被调用方与测试引用）。

## 与另外两个模块的分工

* `sheet_stats.py` 算「一张表逐列的统计」（纯计算，不认识文件）—— 本模块 import 它；
* `excel_view.py` 负责单元格与表头的**排版** —— 本模块只用它的 `_cell`。
  反过来 `excel_view` 不认识工作簿，所以两边不成环。
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Optional, Sequence

from services.ai.budget import truncate_text
from services.ai.excel_view import _cell

# 配表整表统计（列上限 / 分位数 / 去重取值 / 空值率）。它是纯函数，与「向模型交付
# 什么」的其余部分没有耦合 —— 单独一层（见模块 docstring），也好被完整单测。
from services.ai.sheet_stats import _SheetStats
from utils.logger import log_print

# 纯文本内容（Excel 渲染成表）
# --------------------------------------------------------------------------

# 能进 `_read_excel_sheets` 的扩展名：判定依据是**解析器的实际能力**，不是平台的配表清单。
#
# `_read_excel_sheets` 用 openpyxl，只认 OOXML 工作簿。平台的配表清单
# （`ExcelDiffCacheService.is_excel_file` = `.xlsx/.xls/.xlsm/.xlsb/.csv`）比它宽，
# 直接拿来会踩两个坑：
#   * `.csv` / `.tsv` 本来是纯文本，文本分支能把**完整内容**交给模型；若按「配表」进
#     表格分支，`load_workbook` 解析失败 → 模型只收到「内容无法解析成文本表格」，
#     等于把读得到的内容说成读不了；
#   * `.xls` / `.xlsb` 是 OLE/二进制，openpyxl 同样不支持，进文本分支会得到显式的
#     「不是文本也不是配表」说明（不会抛异常）。
# 所以这里只放 openpyxl 真正能打开的那几种。
_WORKBOOK_EXTENSIONS = ('.xlsx', '.xlsm', '.xltx', '.xltm')


def _is_openpyxl_workbook(path: str) -> bool:
    """这个路径该按工作簿解析吗（= openpyxl 读得动吗）。"""
    return Path(str(path or '')).suffix.lower() in _WORKBOOK_EXTENSIONS


def parse_sheet_window(window: str) -> Optional[tuple]:
    """`lines` 在**配表**上表示「第几张工作表」（1 起算的闭区间）。认不出来返回 `None`。

    ## 为什么配表按「张」而不是按「行」

    文本 / 代码的 `lines` 是行号，因为行号是模型在代码里定位的坐标（`@@ -1180,7` 与正文
    行号是同一套）。而配表的正文是**每张表各自的统计 + 前若干行**，根本没有「整个文件的
    第几行」这个概念 —— Agent 端那条路甚至显式标了 `kind: "excel"` 并注明「不要再按行号
    包一层」。所以这里换一个模型说得清、也对得上的坐标：**第几张表**。

    `protocol._normalize_line_window` 已经把 `lines` 规整成 `N` / `N-M` / `""`，这里只做
    「认不认得出」：**认不出就当没给**，与 `utils/content_window.parse_line_window` 同一条
    口径 —— 窗口写坏的代价只能是拿到默认那一段，不能是丢掉整条请求。
    """
    text = str(window or "").strip()
    if not text:
        return None
    match = re.fullmatch(r"(\d{1,7})(?:-(\d{1,7}))?", text)
    if match is None:
        return None
    start = int(match.group(1))
    if start < 1:
        return None
    end = int(match.group(2)) if match.group(2) else start
    if end < start:
        return None
    return start, end


def _row_has_content(cells: Sequence[str]) -> bool:
    """这一行有没有内容（全部是空串 / 空白串就是没有）。

    判据与 `_SheetStats.add_row` 里那句「整行为空」逐字一致：**空白串不算内容**
    （`'   '` 在配表里是占位，不是字段名），免得一个空行被当成列名行。
    """
    return any(str(value or "").strip() for value in cells)


def _read_excel_sheets(
    raw: bytes,
    *,
    max_rows: int,
    path: str = "",
    window: str = "",
    char_budget: int = 0,
    header_rows=None,
    header_name_row=None,
) -> Optional[str]:
    """把 xlsx 的字节渲染成文本表格。**自报家门、自我收敛、按工作表可点名。**

    Excel 的「完整内容」本质就是一张表；按二进制拒绝它会让模型完全看不到这张表长什么样。

    ## 排版：所有工作表的**统计**先发，所有**正文**后发

    截断只砍尾巴（`budget.truncate_text`），所以「先统计后正文」是这个函数的立意 ——
    统计是「值合不合理」这类判断唯一拿得到的比较基准，被砍掉就等于这条路径白跑了。

    但这件事故**不能按工作表各排各的**。原先的排版是「表1 统计 → 表1 正文 → 表2 统计 → …」：
    表 1 的正文一多，后面每张表的统计就全在截断线之外了。线上真实的一本书就是这样 ——
    0_常规属性 3,053 行、整份渲染 25,261 字，截到 11,000 之后「2_M scs属性」的统计一个字
    都没到，而「属性叠加方式=4 是否合理」正需要拿它的统计当基准。

    ## 自我收敛：**不把「砍掉一半」留给预算层**

    上面那个修法解决了统计，但正文仍然可能被整段尾截断 —— 而配表**没有补救路径**：
    `file_content` 的 `lines` 对代码是行窗口、对配表原先被完全忽略，所以模型被告知
    「有 1 条上下文因长度上限被截断」，然后什么都做不了。同轮里另外三个会返回长内容的
    工具（`file_diff` / `read_reference` / `commit_detail`）都有「分段 + 点名」。

    所以这个函数**自己按 `char_budget` 收敛**，并把「一共有几张表 / 这次给了哪张 /
    别的怎么要」写在最前面。它交给上层的文本**永远不会超预算** —— 于是配表不再出现
    「因长度上限被截断」，取而代之的是一句模型能照做的指路：要看第 2 张，写 `"lines": "2"`。

    ## 「交出去的文本不超预算」有一个前提：额度里要留出平台自己会补的那一行

    这个函数收敛到的是**它自己**那份文本的长度。单机部署下它就是最终文本，没问题；
    但多节点部署下平台还会在它前面补一行出处（`_render_agent_file_content`）——
    而那行**不在**这个函数的额度内。原先配表是按满额度渲染的，于是
    「顶满 11,000 + 补 27 字」必然超上限，上层的 `ContextTools` 再砍一刀：
    同一份配表在单机下不报「被截断」、在多节点下报，被砍掉的正是末尾那行
    「正文只展示了前 N 行」与这个函数自己的截断标记。

    修法在 `file_content`：配表那条路**两端都按扣掉那一行之后的额度**渲染
    （见 `_AGENT_EXCEL_NOTE`）—— 两端正文逐字节相同，补上出处之后仍不超上限。

    单张工作表自身超额度时仍会落到下面那句 `truncate_text`（只砍尾巴，重问同一张表
    得到逐字节相同的结果）——那一段是**真的拿不回来**，所以 `engine._truncation_note`
    必须把这件事说给模型听（但只能说「表内被砍掉的行」，不能说成整类配表）。

    ## 参数

    * `max_rows`：每张表的正文最多给多少行（统计不受它影响，永远覆盖整表）。
    * `path`：写进抬头（与 `_render_text_content` 同一个约定）。
    * `window`：`""` = 从第 1 张开始、尽量多给；`"2"` / `"2-3"` = 点名要第几张。
    * `char_budget`：0 = 不限（小工具与既有单测用）；生产路径传单条上限。
    * `header_rows` / `header_name_row`：仓库上的「表头行数」与「名称行」
      （`Repository.header_rows` / `Repository.header_name_row`）。**不传就是今天的行为**
      （列名取第一个非空行，其余行全算数据），所以既有调用方与单测一个字都不用改。

    ## 表头坐标为什么必须由仓库配置说了算（2026-09-20）

    这一层原先**完全不认识**这两个配置：列名一律取「第一个非空行」。而 diff 引擎的列名
    由 `header_name_row` 决定（`services/diff_service.py::_plan_name_row`），表头块由
    `header_rows` 决定。于是「第 1 行是大标题、第 2 行才是字段名」（`header_name_row=2`
    存在的唯一理由）的表上：

    * diff 侧列名取第 2 行，AI 侧却把第 1 行当列名；
    * 第 2 行（真正的字段名行）在 AI 侧被当成**数据第一行**统计进去 —— 每个文本列于是
      多出一个假的「只出现一次的取值」，而模型正是拿这些取值分布判断「这个值合不合理」。

    两端都读同一份仓库配置之后，列名与数据行的坐标就一致了。配了 `header_rows=3` 的表，
    物理行 2、3 也不再进统计与正文（diff 引擎早就是这样）。

    取法与原函数一样只动**坐标**，不改编排：`header_rows` 的那几行既不进统计、也不进
    正文，只在工作表抬头里如实写一句「表头共 N 行不计入」，免得模型以为这几行凭空消失。
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover —— openpyxl 是平台的既有依赖
        return None

    # 规范化只走 diff 引擎那两份实现（`_header_row_count` / `_header_name_row`）：
    # 同一件事不能有第二套「几算合法、越界怎么办」的判断 —— 那正是本次要收敛的东西。
    from services.diff_service import DiffService

    header_count = DiffService._header_row_count(header_rows)
    name_row = DiffService._header_name_row(header_name_row, header_count)

    try:
        book = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 —— 解析失败要变成可读说明，不能炸掉整轮
        log_print(f"⚠️ AI 取数：Excel 内容解析失败 {type(exc).__name__}: {exc}")
        return None

    # 逐表收集：`(表名, 统计行, 正文行, 数据总行数)`。正文只收 `max_rows` 行 —— 统计已经在
    # `stats.add_row` 里过了全部行，正文收多了没用（行窗口不存在），却要为此把几千行
    # 字符串留在内存里。
    collected: list[tuple[str, list[str], list[str], int]] = []
    try:
        for name in book.sheetnames:
            sheet = book[name]
            stats = _SheetStats()
            labels: list[str] = []
            first_row: list[str] = []
            body: list[str] = []
            for physical, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                cells = ["" if value is None else str(value) for value in (row or ())]
                if physical == 1:
                    first_row = cells
                if name_row > 1 and physical == name_row:
                    # 配置了名称行：列名**就是这一行**（与 diff 引擎的 `_plan_name_row` 同口径）。
                    labels = cells
                    continue
                if name_row == 1 and not labels and _row_has_content(cells):
                    # 未配置名称行（或配置为 1）＝今天的行为：第一个非空行按列名处理。
                    # 它只当标签，不进统计（表头文字当成取值会让每个文本列都多出一个
                    # 「只出现一次」的值）。
                    labels = cells
                    continue
                if physical <= header_count:
                    # 表头块（物理行 1..header_count）里剩下的行：既不是数据，也不是列名。
                    # 不进统计、不进正文 —— 与 diff 引擎「只有物理行 > header_rows 才算
                    # 数据行」同一口径。
                    continue
                stats.add_row(row)
                if len(body) >= max_rows:
                    # 统计要覆盖整表，所以这里不能 break：只停止累积正文。
                    continue
                if row is None or all(value is None for value in row):
                    continue
                body.append("- " + " ｜ ".join(_cell(value) for value in row))
            if not labels and name_row > 1 and _row_has_content(first_row):
                # 配置了名称行、但这张表短到没有那一行：与 diff 引擎同一口径 ——
                # `_plan_name_row` 取不到名称行时「不改名」，列名仍是 `header=0` 读到的第 1 行。
                labels = first_row
            collected.append(
                (
                    str(name),
                    stats.render(labels, header_count=header_count, name_row=name_row),
                    body,
                    stats.rows,
                )
            )
    finally:
        try:
            book.close()
        except Exception:  # noqa: BLE001
            pass

    return _assemble_workbook(
        collected, path=path, window=window, char_budget=char_budget, max_rows=max_rows
    )


# 抬头与「没给的那几张表怎么要」那两行必须留下来的预留：前者是「这是哪一段」的唯一出处，
# 后者是模型唯一能照做的补救动作（见 `_read_excel_sheets`）。
_EXCEL_HEADER_RESERVE = 700
# 统计那一档最多吃掉预算的多大比例。它是**基准**，比正文值钱，所以先给它一份保底额度；
# 但也不能全给它 —— 一张 24 列的宽表统计就要 2,000 字上下，几张表就能把预算吃光。
_EXCEL_STATS_SHARE = 0.5
# 一张表在正文段里**至少**要给到几行，否则这一块不值得占额度（统计本来就覆盖整表了）。
# 低于它就不给这一张的正文，改为在抬头里点名 —— 模型写一句 `"lines": "3"` 就能单独拿到，
# 而「每张表各给一行」既占额度又几乎不含信息。
_EXCEL_MIN_BODY_ROWS = 20


def _block_chars(lines: Sequence[str]) -> int:
    """一组行的字符数（含换行）—— 排版的额度判断只用它。"""
    return sum(len(line) + 1 for line in lines)


def _assemble_workbook(
    collected: Sequence[tuple],
    *,
    path: str,
    window: str,
    char_budget: int,
    max_rows: int,
) -> str:
    """把逐表收集到的内容排成「统计 → 正文」两段，并在预算内**说清缺了什么**。

    单独提出来是因为它是这个渲染里唯一会算错的地方（谁进得去、谁进不去、进去了几张），
    而它只依赖收集结果与一个字符数 —— 可以直接喂数进来验，不必造真 xlsx。
    """
    total_sheets = len(collected)
    limit = max(0, int(char_budget or 0))
    # 统计档的额度：预算的一半（至少留 2,000 —— 低于它连一段真实统计都放不下，那这一档
    # 就没有存在的意义了），并且**扣掉抬头那 700 的预留**：统计把预算吃光、抬头被截掉，
    # 模型就既不知道这是哪份文件、也不知道怎么要剩下的。
    stats_room = 0
    if limit:
        stats_room = max(
            0, min(limit - _EXCEL_HEADER_RESERVE, max(2_000, int(limit * _EXCEL_STATS_SHARE)))
        )

    stats_lines: list[str] = []
    stats_skipped: list[str] = []
    for index, (name, stats, _body, _rows) in enumerate(collected, start=1):
        # 空表（一行数据都没有）`render` 返回空列表 —— 不冒出一段「按 0 行算出」的统计。
        if not stats:
            continue
        block = [f"#### 工作表「{name}」", *stats]
        # 第一张的统计永远给（`stats_lines` 为空时不拦）：一张统计都没有的话，这一档就白设了。
        if limit and stats_lines and _block_chars(stats_lines) + _block_chars(block) > stats_room:
            stats_skipped.append(f"第 {index} 张「{name}」")
            continue
        stats_lines.extend(block)

    lines: list[str] = []
    if stats_lines:
        lines.append(
            "### 整表统计（覆盖每一张工作表的**全部**行；与下面只展示一部分正文无关。"
            "判断某个值是否合理时用它做比较基准）"
        )
        lines.extend(stats_lines)
        if stats_skipped:
            lines.append(
                f"- （另有 {len(stats_skipped)} 张表的统计因篇幅没列：{'、'.join(stats_skipped)}。"
                "它们的正文可以按下面的方法单独索取。）"
            )
        lines.append("")

    # 正文档：点名了就给点名的那些，没点名就从第 1 张开始、尽量多给。
    span = parse_sheet_window(window)
    asked = span is not None
    wanted = (
        list(range(span[0], min(span[1], total_sheets) + 1))
        if span is not None
        else list(range(1, total_sheets + 1))
    )

    # 正文段：每张表一个块。`built` 只装**进得去**的那些，兜底时从尾部整块退。
    #
    # 额度预算里先扣掉 `_EXCEL_HEADER_RESERVE`（抬头 + 那句指路的预留），于是「装不下」
    # 只可能发生在正文段 —— 抬头与指路是整个函数唯一能自我补救的地方，它们必须留下。
    used = _block_chars(lines) + (_EXCEL_HEADER_RESERVE if limit else 0)
    built: list[tuple[int, list[str]]] = []
    for index in wanted:
        name, _stats, rows, total_rows = collected[index - 1]

        def _body_block(take: int, name: str = name, rows: list = rows, total_rows: int = total_rows):
            """这个工作表在「只给前 `take` 行」时的正文块。"""
            chunk = [f"### 工作表「{name}」（数据第 1–{take} 行，共 {total_rows} 行）"]
            chunk.extend(rows[:take])
            if not take:
                chunk.append("- （这张表没有数据行。）")
            elif take < total_rows:
                chunk.append(
                    f"- （正文只展示了前 {take} 行，共 {total_rows} 行；"
                    "上面的整表统计覆盖整表。要看**改动的行**请用 `file_diff`，它按改动行给。）"
                )
            return chunk

        # 这张表的正文**给多少行**由额度说了算，`max_rows` 只是上限之一。
        #
        # 这一步不能省：一行宽表（24 列、每列几十字）就要 200 字上下，`max_rows=120` 的
        # 一块就是两万多字 —— 超过了整条上限。只做「整块进不进得去」的判断，结果是第一块
        # 就把预算吃穿，然后退回 `truncate_text` 砍尾巴 —— 正是这次要消掉的那条死路。
        #
        # 逐次减行而不是一次算出「能放几行」：块里还有抬头与那句「只展示了前 N 行」的说明，
        # 它们随 `take` 变，一次算不准；这里按四分之一收敛，几十行以内必然落到位。
        take = len(rows)
        floor = min(len(rows), _EXCEL_MIN_BODY_ROWS)
        block = _body_block(take)
        while limit and take > floor and used + _block_chars(block) > limit:
            take = max(floor, take - max(1, take // 4))
            block = _body_block(take)
        # 第一块永远给（哪怕只是一行）——「一张表都没给」比「给了一张、其余指路」糟得多。
        # 后面的表如果连 `_EXCEL_MIN_BODY_ROWS` 行都放不下，就不占这份额度：它已经被抬头
        # 点名了，模型写一句 `"lines": "3"` 就能单独拿到它。
        if limit and built and used + _block_chars(block) > limit:
            break
        built.append((index, block))
        used += _block_chars(block) + 1  # +1 是块之间那个空行

    def _render(blocks: Sequence[tuple[int, list[str]]]) -> str:
        """把「抬头 + 统计档 + 正文档」拼成最终文本。抬头里的那张清单由 `blocks` 决定 ——
        兜底退块之后它必须跟着改，否则抬头会说「给了第 3 张」而正文里没有。"""
        shown = [index for index, _chunk in blocks]
        head = [f"配表正文：{path}" if path else "配表正文"]
        head[0] += f"（共 {total_sheets} 张工作表）"
        if shown:
            head[0] += f"；本次给了第 {'、'.join(str(item) for item in shown)} 张的正文"
        elif asked:
            head[0] += (
                f"；你要的第 {span[0]}-{span[1]} 张不存在（这份工作簿一共只有 {total_sheets} 张）"
            )
        else:
            head[0] += "；本次没有给任何一张表的正文"
        missing = [index for index in range(1, total_sheets + 1) if index not in shown]
        if missing:
            example = f"{missing[0]}-{missing[0] + 1}" if len(missing) > 1 else str(missing[0])
            head.append(
                f'要看别的表，在请求里写 "lines": "<第几张表>"（例如 "{example}"）。'
                f"另有 {len(missing)} 张表的正文没给："
                + "、".join(f"第 {index} 张「{collected[index - 1][0]}」" for index in missing)
            )
        elif asked:
            head[0] += "（你点名要的那几张）"
        chunks = [*head, "", *lines]
        if blocks:
            chunks.append(f"### 工作表正文（每张表最多展示前 {max_rows} 行）")
            for _index, block in blocks:
                chunks.extend(block)
                chunks.append("")
        return "\n".join(chunks).rstrip() + "\n"

    text = _render(built)
    # 兜底：预估偏小时**整块往回退**，而不是整段 `truncate_text` —— 后者会砍掉最后那句
    # 「要看别的表怎么写」，那正是配表唯一的重来路径（也正是这次改动要解决的事）。
    while limit and len(text) > limit and len(built) > 1:
        built = built[:-1]
        text = _render(built)
    if limit and len(text) > limit:
        # 只剩一块了还是超（一张表本身就有几万字）：这时只能收它的尾巴，抬头照旧。
        text = truncate_text(text, limit)[0]
    return text
