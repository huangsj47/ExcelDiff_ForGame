# -*- coding: utf-8 -*-
"""AI 取数：`file_content` 的「配表」分支必须真的能跑（services/ai/platform_provider.py）

## 为什么需要这个文件

`PlatformContextProvider.file_content` 用「这个路径是不是配表」来决定把内容渲染成
文本表格还是直接解码成文本：

    from services.excel_cache_service import is_excel_file     # ← 这个模块根本不存在
    if is_excel_file(path):
        rendered = _read_excel_sheets(raw, max_rows=self._max_rows)
        ...

`services/excel_cache_service.py` **不在仓库里**，而 `is_excel_file` 在
`ExcelDiffCacheService` / `WeeklyExcelCacheService` 上都是**实例方法**，
不是模块级函数。于是这条 import 必然抛 `ModuleNotFoundError`，而且它在
`try` 之外（那个 try 只包住取 Git 内容），所以：

* 走不到 `_read_excel_sheets`（真 xlsx 的内容永远拿不到）；
* 抛出的是异常而不是可读说明 —— 与本文件所在模块的契约相反
  （「拿不到」要明说，不能让调用方以为「没有内容」）。

## 分支判定为什么按 openpyxl 的实际能力选（而不是平台的「配表」清单）

`_read_excel_sheets` 用 `openpyxl.load_workbook`，**只认 OOXML 工作簿**
（`.xlsx` / `.xlsm` / `.xltx` / `.xltm`）。所以本分支的判定必须是「openpyxl 读得动吗」：

* 平台的配表清单（`ExcelDiffCacheService.is_excel_file` = `.xlsx/.xls/.xlsm/.xlsb/.csv`）
  比 openpyxl 的能力**宽**。其中 `.csv` 是会踩坑的那个：CSV 本来是纯文本，
  下面的文本分支能原样交给模型完整内容；一旦按「配表」进表格分支，
  `load_workbook` 解析失败 → 模型只会收到「内容无法解析成文本表格」，
  **反而丢掉了本来读得到的内容**（等于把「读到了」说成「读不了」）。
  `.tsv` 同理（且它连平台的配表清单都不在）。
* `.xls` / `.xlsb` 是 OLE/二进制格式，openpyxl 也不支持；它们进文本分支会得到
  下面那条「不是文本也不是配表」的显式说明，不会抛异常。

回归保护：`tests/test_ai_platform_provider.py` 覆盖的是渲染层（`render_diff_payload`），
本文件覆盖的是取数层的这个分支。
"""
import io
import os
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from services.ai.platform_provider import (  # noqa: E402
    PlatformContextProvider,
    _is_openpyxl_workbook,
)
from services.ai.skill_loader import LoadedSkills, SkillDocument  # noqa: E402

COMMIT = 'a' * 40


class TestWorkbookPredicate:
    """判定必须跟 `_read_excel_sheets`（openpyxl）的实际能力对齐。

    锁定这条判断，避免有人「顺手统一成平台的配表清单」而让 CSV/TSV 丢掉可读内容。
    """

    @pytest.mark.parametrize('path', [
        'config/道具表.xlsx', 'config/道具表.XLSX', 'config/道具表.xlsm',
        'config/道具表.xltx', 'config/道具表.xltm',
    ])
    def test_openpyxl_readable_paths_are_workbooks(self, path):
        assert _is_openpyxl_workbook(path) is True, path

    @pytest.mark.parametrize('path', [
        'config/道具表.csv', 'config/道具表.tsv', 'config/道具表.xls',
        'config/道具表.xlsb', 'services/foo.py', 'assets/blob.bin', '', 'noext',
    ])
    def test_everything_else_is_not(self, path):
        assert _is_openpyxl_workbook(path) is False, (
            f'{path!r} 被判成工作簿了。\n'
            f'`.csv`/`.tsv` 走文本分支才能把完整内容交给模型，'
            f'`.xls`/`.xlsb` openpyxl 根本打不开。'
        )


@pytest.fixture()
def provider():
    """只需要能构造出来的 provider —— `_commit_row` / 取内容都会被替换。"""
    doc = SkillDocument(
        name='SKILL.md', description='', path=Path('SKILL.md'), text='平台协议',
        content_hash='h',
    )
    return PlatformContextProvider(
        loaded=LoadedSkills(
            platform_skill=doc,
            platform_references=(),
            project_manifest=None,
            project_references=(),
            project_skills=(),
            readable={},
            project_slug=None,
            revision='rev',
        )
    )


@pytest.fixture()
def fetch(monkeypatch):
    """把「查提交行 + 从 git 取内容」替换成内存里的字节，只留被测分支。"""
    def _install(content: bytes):
        monkeypatch.setattr(
            PlatformContextProvider, '_commit_row',
            lambda self, commit, path: SimpleNamespace(commit_id=commit, path=path, repository=object()),
        )

        def _get_content(repository, commit_id, path):
            return content

        import services.vcs_content_service as vcs

        monkeypatch.setattr(vcs, 'get_file_content_from_git', _get_content, raising=False)
    return _install


@pytest.fixture(scope='module')
def repository_id():
    """Agent 侧那条路要按 id 查真实的仓库行（`read_file_content_for_agent` 会查库）。

    作用域是 module：`create_tables()` 每次都会把整张表清单打进日志；项目 code 带 uuid
    后缀，所以不会撞会话级共用测试库里的唯一约束。
    """
    import uuid

    from app import app, create_tables, db
    from models import Project, Repository

    token = uuid.uuid4().hex[:10]
    with app.app_context():
        create_tables()
        project = Project(code=f'XB{token}', name=f'agent-excel-branch-{token}')
        db.session.add(project)
        db.session.flush()
        repository = Repository(
            project_id=project.id, name=f'repo-xb-{token}', type='git',
            url='https://example.invalid/game.git', branch='main', clone_status='completed',
        )
        db.session.add(repository)
        db.session.commit()
        return repository.id


def _xlsx_bytes(rows, sheet_name='道具表'):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=col_index, value=value)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


class TestExcelBranchIsReachable:
    """真 xlsx：必须渲染成文本表格，而不是抛异常、也不是「无法解析」。"""

    def test_real_xlsx_returns_rendered_table(self, provider, fetch):
        fetch(_xlsx_bytes([['ID', '攻击'], ['1001', '30']]))
        got = provider.file_content(COMMIT, 'config/道具表.xlsx')
        assert isinstance(got, str), f'返回的不是文本：{got!r}'
        assert '道具表' in got, f'没有渲染工作表名：{got!r}'
        assert '1001' in got and '攻击' in got, f'没有渲染单元格内容：{got!r}'
        assert '无法解析' not in got, f'真 xlsx 被说成解析不了：{got!r}'

    def test_xlsm_extension_also_renders(self, provider, fetch):
        """`.xlsm` 也是 openpyxl 能读的 OOXML 工作簿。"""
        fetch(_xlsx_bytes([['ID'], ['1001']]))
        got = provider.file_content(COMMIT, 'config/道具表.xlsm')
        assert isinstance(got, str) and '1001' in got, got

    def test_parse_failure_of_a_workbook_is_explicit(self, provider, fetch):
        """扩展名像工作簿但内容坏掉时：给可读说明，**不是**异常，也不是「没有内容」。"""
        fetch(b'not a workbook at all')
        got = provider.file_content(COMMIT, 'config/道具表.xlsx')
        assert isinstance(got, str), f'抛异常/返回 None 都不允许：{got!r}'
        assert got != '', '空串在这个契约里是「确实没有内容」，不能说谎'
        assert '无法解析' in got, f'解析失败必须明说：{got!r}'


class TestNonWorkbookPathsGoToText:
    """`.csv` / `.tsv` / 普通文本：必须**逐行原样**给模型，不能进表格分支。

    「原样」的判定标准是**内容行一字不改地出现在返回里**（带行号，见
    `platform_provider._render_text_content`），不是「返回与文件字节逐字相同」——
    后者会把抬头与行号一起钉死，而那两样正是模型据以定位的证据。
    """

    def _assert_lines_are_intact(self, got, lines):
        """每一行都必须以 `行号│原文` 的形态出现，且顺序不变。"""
        assert isinstance(got, str) and got, f'不允许返回 None/空串：{got!r}'
        for number, text in enumerate(lines, start=1):
            assert f'{number}│{text}' in got, (
                f'第 {number} 行没按「行号│原文」给出来：{text!r}\n完整返回：{got!r}'
            )

    def test_csv_is_returned_as_text(self, provider, fetch):
        fetch('id,name\n1,Alice\n'.encode('utf-8'))
        got = provider.file_content(COMMIT, 'config/道具表.csv')
        assert '无法解析' not in got, (
            f'CSV 被当成工作簿解析了（模型会收到「无法解析」而不是这张表）：{got!r}'
        )
        self._assert_lines_are_intact(got, ['id,name', '1,Alice'])

    def test_tsv_is_returned_as_text(self, provider, fetch):
        fetch('id\tname\n1\tAlice\n'.encode('utf-8'))
        got = provider.file_content(COMMIT, 'config/道具表.tsv')
        assert '无法解析' not in got, f'TSV 被当成工作簿解析了：{got!r}'
        self._assert_lines_are_intact(got, ['id\tname', '1\tAlice'])

    def test_plain_python_file_is_returned_as_text(self, provider, fetch):
        fetch('def f():\n    return 1\n'.encode('utf-8'))
        got = provider.file_content(COMMIT, 'services/foo.py')
        self._assert_lines_are_intact(got, ['def f():', '    return 1'])

    def test_the_window_says_which_lines_it_is_and_how_many_there_are(
        self, provider, fetch
    ):
        """抬头是模型写结论时的坐标系（「第 1180 行」），行号必须从 1 起、总数要对。"""
        fetch(''.join(f'line {i}\n' for i in range(1, 11)).encode('utf-8'))
        got = provider.file_content(COMMIT, 'services/foo.py')
        assert '共 10 行' in got, f'没告诉模型这个文件有多少行：{got!r}'
        assert '第 1–10 行' in got, f'没告诉模型这是哪一段：{got!r}'

    def test_binary_junk_says_it_cannot_be_shown(self, provider, fetch):
        """非文本非配表：明说无法展示，且**不**说成「没有内容」。"""
        fetch(b'\xff\xfe\x00\x01\x02')
        got = provider.file_content(COMMIT, 'assets/blob.bin')
        assert isinstance(got, str) and got != '', f'不允许返回 None/空串：{got!r}'
        assert '没有内容' in got or '无法展示' in got, got


class TestContentEmptyContract:
    """既有契约保持不变：取到空内容 → 空串；取不到 → None。"""

    def test_empty_content_is_empty_string(self, provider, fetch):
        fetch(b'')
        assert provider.file_content(COMMIT, 'config/道具表.xlsx') == ''

    def test_missing_commit_row_is_none(self, provider, monkeypatch):
        monkeypatch.setattr(PlatformContextProvider, '_commit_row', lambda self, c, p: None)
        assert provider.file_content(COMMIT, 'config/道具表.xlsx') is None


# ---------------------------------------------------------------------------
# Agent 侧的同一条分支（2026-09-19 补）
# ---------------------------------------------------------------------------
#
# platform/agent 模式下代码文件与配表的正文都只有业务节点读得到，而
# `services/agent_file_content_reader.py` 原先**一律按 UTF-8 解码**（`errors='replace'`）——
# xlsx 是 ZIP，解码出来是一段二进制乱码，而抬头还写着「共 N 行；下面是第 a–b 行」。
# 模型据此写的结论全是错的，**比「读不到」更糟**：它不认为自己拿到的不是内容。
#
# 所以两端必须是同一个渲染函数（这里是 `_read_excel_sheets`），渲染出来的文本也必须是
# 同一种形态（工作表名 + 整表统计 + 前若干行）—— 否则同一次索取在单机与多节点下会给出
# 不同的内容，而「同一份规则只有一份实现」正是 `utils/content_window` 那条纪律。


class TestAgentSideWorkbookBranch:
    """`services/agent_file_content_reader.read_file_content_for_agent` 的配表分支。"""

    def _read(self, monkeypatch, repository_id, path, content, payload=None):
        from app import app

        with app.app_context():
            from services.agent_file_content_reader import read_file_content_for_agent

            monkeypatch.setattr(
                'services.vcs_content_service.get_file_content_from_git',
                lambda repository, commit_id, file_path: content,
            )
            return read_file_content_for_agent({
                'repository_id': repository_id, 'commit_id': COMMIT, 'file_path': path,
                **(payload or {}),
            })

    def test_a_workbook_is_rendered_as_a_table_not_decoded_as_bytes(self, monkeypatch, repository_id):
        """**这一条就是「模型拿到二进制乱码却以为是内容」的回归。**"""
        got = self._read(monkeypatch, repository_id, 'config/道具表.xlsx',
                         _xlsx_bytes([['ID', '攻击'], ['1001', '30']]))

        assert got['kind'] == 'excel', f'平台侧要靠它决定不套行号：{got}'
        assert '道具表' in got['content'] and '1001' in got['content'], got
        assert '整表统计' in got['content'], '形态必须与平台本地那条路一致'
        # 反面：按 UTF-8 解码出来的乱码里必然有替换字符或 PK 头。
        assert '�' not in got['content'], '拿到的还是二进制乱码'
        assert 'PK' not in got['content'][:64], got['content'][:64]

    def test_the_row_limit_is_honoured_and_defaults_to_the_platform_value(self, monkeypatch, repository_id):
        """`max_rows` 由平台传（与平台本地同一个默认值），传了就按它渲染。"""
        rows = [['ID', '价值']] + [[str(i), str(i)] for i in range(1, 60)]

        default_rows = self._read(monkeypatch, repository_id, 'config/道具表.xlsx', _xlsx_bytes(rows))
        small = self._read(monkeypatch, repository_id, 'config/道具表.xlsx',
                           _xlsx_bytes(rows), {'max_rows': 3})

        # 判据是**渲染出来的行数**：限 3 行那一份必须明显更短（不是「某个字符串在不在」）。
        assert len(small['content'].split('\n')) < len(default_rows['content'].split('\n')), (
            f'行数上限没生效：{small["content"]!r}'
        )
        assert 'ID' in small['content'], small  # 表头那几行仍要在（它决定这张表是什么）

    def test_a_broken_workbook_raises_instead_of_returning_junk(self, monkeypatch, repository_id):
        """扩展名像工作簿但内容坏了：抛（→ 平台把原因写给模型），**不能**返回乱码。"""
        with pytest.raises(Exception) as excinfo:
            self._read(monkeypatch, repository_id, 'config/道具表.xlsx', b'not a workbook at all')
        assert '无法解析' in str(excinfo.value), str(excinfo.value)

    def test_a_csv_is_still_text(self, monkeypatch, repository_id):
        """`.csv`/`.tsv` 不进表格分支（openpyxl 打不开它们，而文本分支能原样给出内容）。"""
        got = self._read(monkeypatch, repository_id, 'config/道具表.csv',
                         'id,name\n1,Alice\n'.encode('utf-8'))

        assert got.get('kind') != 'excel', got
        assert got['content'].split('\n') == ['id,name', '1,Alice'], got
        assert got['total_lines'] == 2, got

    def test_a_binary_file_gets_the_same_notice_as_the_platform_path(self, monkeypatch, repository_id):
        """非工作簿的二进制（`.bin`）：**与平台本地那条路给出同一句话**。

        修前这里与平台本地是两份不同的文本：平台侧严格 `utf-8` 解码失败 → 回一句
        「[无法展示的内容] …不是文本…」，Agent 侧 `errors='replace'` → 回一段带替换符的
        乱码（本用例原先断的就是这个：`total_lines >= 1`）。同一次索取在单机与多节点下
        于是给出两份不同的文本，而模型据此写出的结论也会不同。

        现在两端都走 `utils.text_decoding`：真正的二进制（魔数 / NUL）给同一句话
        （`binary_content_notice`），并且用 `kind: "binary"` 告诉平台侧不要再按行号包一层。
        """
        got = self._read(monkeypatch, repository_id, 'assets/blob.bin', b'\x00\x01\x02\x03\xff')

        assert got.get('kind') == 'binary', got
        assert '无法展示' in got['content'] and '没有内容' in got['content'], got
        assert 'total_lines' not in got, got


class TestPlatformSideRendersTheAgentWorkbook:
    """平台侧收到 `kind == "excel"` 时：直接用，**不能再按行号包一层**。

    Agent 回的正文已经是「整表统计 + 前若干行」的形态（与平台本地同一个渲染函数），
    没有「第 a–b 行」这个概念。这里再套一层行号会变成「行号套行号」，模型会把每一行的
    工作表正文当成文件行号去引用 —— 一个它自己发现不了的坐标错误。
    """

    def test_the_excel_kind_is_passed_through_with_its_own_header(self):
        from services.ai.platform_provider import _render_agent_file_content

        text = _render_agent_file_content(
            {'kind': 'excel', 'file_path': 'config/道具表.xlsx',
             'content': '配表正文：config/道具表.xlsx（共 1 张工作表）；本次给了第 1 张的正文\n'
                        '### 工作表正文（每张表最多展示前 120 行）\n'
                        '### 工作表「道具表」（数据第 1–2 行，共 2 行）\n- 1001 ｜ 30\n'},
            path='config/道具表.xlsx',
        )

        assert '配表正文' in text, text
        assert '工作表「道具表」' in text, text
        assert '1│1│' not in text, f'套了两层行号：{text}'
        assert '共 3000 行' not in text, text
        # **抬头只有一个**：它由渲染函数自己写（只有渲染函数说得清「共几张表 / 给了哪几张 /
        # 别的怎么要」），平台侧再补一个就成了两行出处、两个口径。
        assert text.count('配表正文：') == 1, text
        # 出处仍然要在 —— 这一句是模型判断「这份正文是谁读的」的唯一依据。
        assert '业务节点（Agent）' in text, text

    def test_an_empty_workbook_answer_is_not_read_as_no_content(self):
        from services.ai.platform_provider import _render_agent_file_content

        text = _render_agent_file_content(
            {'kind': 'excel', 'file_path': 'config/道具表.xlsx', 'content': ''},
            path='config/道具表.xlsx',
        )

        assert '无法解析' in text, text
        assert '不等于「没有内容」' in text, text

    def test_a_plain_text_answer_keeps_the_line_numbering(self):
        """反面：`kind` 不是 excel 时，行号那一套照旧（不能把两条路混起来）。"""
        from services.ai.platform_provider import _render_agent_file_content

        text = _render_agent_file_content(
            {'content': 'line 1\nline 2', 'start_line': 1, 'end_line': 2, 'total_lines': 2},
            path='src/fight.lua',
        )

        assert '1│line 1' in text and '2│line 2' in text, text
