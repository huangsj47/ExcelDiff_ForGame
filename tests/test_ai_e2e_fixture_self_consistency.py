# -*- coding: utf-8 -*-
"""e2e 造数脚本必须**自洽**：那份「干净的小 delta」不许悄悄改出别的东西。

## 这一条钉的是哪一次真实污染

`scripts/e2e_local_repo/phases.py incremental` 的第二步是「只改一个文件（铁剑价格），
再看增量做差给了什么」。但 `make_fixture.round3()` 原先拿 `ITEMS_V2` 当基，而
`ITEMS_V2` 里有两行 `build()` 从没写进去过的物品 —— 1006（高级药水）与 1007（秘银剑）。
`round2()` 全仓**从来没有被调用过**（`phases.py:192` 只调 `round3`），所以在一个
`build()` 出来的仓库上跑 `incremental`，这笔「干净的小 delta」会**新增那两行**。

后果不是造数难看，是**报告被写坏**：实测 run 52 / run 53 的报告把 1006/1007 写成
「本期删除」的高风险，而真实 Git 历史里这两个 ID 从未存在过 —— 核对仓库三个提交的原表
就能证伪，属于最伤信任的那一类误报。用户的原话是「这两个 ID 从未存在」。

## 判据为什么落在**行为**上

「`round3` 用的是 `ITEMS_AT_BUILD_TIP` 而不是 `ITEMS_V2`」这句话用静态断言钉不住 ——
它只证明某一行怎么写，不证明造出来的仓库里有什么。所以这里**真的把仓库造一遍**
（`build()` + `round3()`，都在临时目录里），再逐个提交读那份 `物品表.xlsx`：

* 每一个提交里都只能有 1001–1005；
* `round3` 那一笔只能动一个文件，且改动只是铁剑的价格。

`ITEMS_V2` 仍然保留 1006/1007（`round2()` 要用它验「新增两行」那条路）——**它本身没错，
错的是拿它当 `round3` 的基**。所以这里也不要求删掉它。
"""
from __future__ import annotations

import io
import subprocess

import pytest

from scripts.e2e_local_repo import make_fixture as mf

PHANTOM_IDS = {1006, 1007}


def _isolate(monkeypatch, tmp_path):
    """把造数过程整个挪到临时目录 —— **不许碰本地那份 e2e 仓库**。

    本地 `.pytest_tmp/e2e` 是联调用的实物（平台的仓库 id=3 指着它）。跑一次测试就把它
    删掉重建，正在用的人会莫名其妙地掉数据。
    """
    src = tmp_path / "gitsrc"
    origin = tmp_path / "origin.git"
    monkeypatch.setattr(mf, "SRC", src)
    monkeypatch.setattr(mf, "ORIGIN", origin)
    monkeypatch.setattr(mf, "BASE", tmp_path)
    return src


def _git(src, *args):
    return subprocess.run(["git", *args], cwd=str(src), capture_output=True, text=True)


def _item_ids(src, rev: str) -> list:
    """某个提交上 `config/物品表.xlsx` 里的 ID（从 git 对象直接读，不动工作区）。"""
    from openpyxl import load_workbook

    blob = subprocess.run(
        ["git", "show", f"{rev}:config/物品表.xlsx"], cwd=str(src), capture_output=True
    ).stdout
    assert blob, "这个提交里没有 config/物品表.xlsx"
    book = load_workbook(io.BytesIO(blob))
    try:
        sheet = book.active
        return [row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    finally:
        book.close()


def _price_rows(src, rev: str) -> dict:
    """某个提交上物品表的 `{id: 价格}`（比整表重写与「只改一格」就靠它分辨）。"""
    from openpyxl import load_workbook

    blob = subprocess.run(
        ["git", "show", f"{rev}:config/物品表.xlsx"], cwd=str(src), capture_output=True
    ).stdout
    book = load_workbook(io.BytesIO(blob))
    try:
        return {
            row[0]: row[3]
            for row in book.active.iter_rows(min_row=2, values_only=True)
            if row[0]
        }
    finally:
        book.close()


def test_the_incremental_commit_adds_nothing_but_the_price_change(monkeypatch, tmp_path):
    """`build()` + `round3()` 之后：**每个提交都只有 1001–1005**，最后一笔只改一格。

    这一条同时钉住三件事，缺一不可：

    1. 没有任何提交含 1006/1007（污染源就是这里）；
    2. `round3` 那一笔**只动一个文件**（「干净的小 delta」这句话的判据）；
    3. 它动的是铁剑的价格 —— 与它的提交信息一致（不是「顺手把整张表重写了一遍」）。
    """
    src = _isolate(monkeypatch, tmp_path)

    mf.build()
    before = _git(src, "rev-parse", "master").stdout.strip()
    mf.round3()
    after = _git(src, "rev-parse", "master").stdout.strip()

    assert after != before, "round3 没有产生新提交"

    revs = _git(src, "rev-list", "master").stdout.split()
    assert len(revs) == 4, "build 三个提交 + round3 一个"
    for rev in revs:
        ids = _item_ids(src, rev)
        assert not (set(ids) & PHANTOM_IDS), (
            f"提交 {rev} 里出现了 {sorted(set(ids) & PHANTOM_IDS)} —— "
            "这两个 ID 在真实历史里从未存在过，报告会把它读成「本期删除」"
        )
        assert ids == [1001, 1002, 1003, 1004, 1005], (rev, ids)

    # `-c core.quotepath=false`：默认配置下 git 会把非 ASCII 路径写成
    # `"config/\347\211\251..."` 这样的八进制转义，断言「改的是物品表」就对不上了。
    changed = [
        line
        for line in _git(
            src, "-c", "core.quotepath=false", "diff", "--name-only", before, after
        ).stdout.splitlines()
        if line
    ]
    assert len(changed) == 1, changed
    assert "物品表" in changed[0], changed
    # **只有铁剑那一格变了。** 只断言「一个文件变了」是不够的：整张表重写一遍（或者
    # 顺手多改一行）也是一个文件，而那已经不是「干净的小 delta」了 —— 增量做差要验的
    # 正是「这一格」，多出来的每一格都会让那一次的结论不能归因。
    before_rows, after_rows = _price_rows(src, before), _price_rows(src, after)
    differing = {key for key in before_rows if before_rows[key] != after_rows.get(key)}
    assert differing == {1003}, (differing, before_rows, after_rows)
    assert before_rows[1003] == 200 and after_rows[1003] == 260
    assert set(before_rows) == set(after_rows), "行数变了 —— 那不是「只改价格」"


def test_round2_still_adds_the_two_rows_it_is_meant_to_add(monkeypatch, tmp_path):
    """`round2` 本身就是「新增两行」那条路 —— 它不许被顺手改哑。

    `round3` 修好之后容易生出一种「顺手把 `ITEMS_V2` 也删干净」的改法：那样一来
    「新增物品」这条路就没人验了，而它同样是真实场景。所以这里反向钉一次：
    **`round2` 造出来的仓库里必须有 1006/1007**。

    这条同时说明了两条轮次为什么是**互斥的续集**：`round2` 加进去的正是 `round3`
    的基里没有的那两行 —— 串着跑，`round3` 就会把它们删回去。
    """
    src = _isolate(monkeypatch, tmp_path)

    mf.build()
    before = _git(src, "rev-parse", "master").stdout.strip()
    mf.round2()
    after = _git(src, "rev-parse", "master").stdout.strip()

    assert _item_ids(src, before) == [1001, 1002, 1003, 1004, 1005]
    assert set(_item_ids(src, after)) >= PHANTOM_IDS, (
        "round2 没加上 1006/1007 —— 要么它被改哑了，要么这条用例的前提不再成立"
    )


@pytest.mark.parametrize(
    "name", ["build", "round2", "round3", "round4", "round5", "round6"]
)
def test_every_entry_point_is_still_callable(name):
    """各入口都得在（`__main__` 那张分发表按名字取用）。"""
    assert callable(getattr(mf, name))


def test_building_twice_gives_the_same_commits(monkeypatch, tmp_path):
    """`build()` 必须**逐字节可复现**：同一个脚本跑两次要给同一批提交哈希。

    ## 为什么这是判据而不是洁癖

    平台里记的全是**提交哈希**：`Repository.last_synced_tip`、`ai_analysis_run` 的做差基线、
    `commits_log` 的每一行。造数只要每次重来都换一批哈希，重建一次就让这些身份**全部失效**
    —— 而症状是「同步成功，增量分析却对不上」这类看不出因果的怪事（实测踩到过：连跑两次
    `build()` 给出 c3 = `2caae19` 与 `1d219e0`）。

    起因是 openpyxl 把**当下时刻**写进 xlsx 字节：`docProps/core.xml` 的
    `dcterms:created/modified`，以及 zip 每个条目头的写入时刻。三处都在文件字节里，
    而 git 认的就是字节 —— 一个字节变，从引入该文件的那个提交起，后面**每一条**提交的哈希全变。

    ## 判据为什么是「两次跑出来的哈希列表相等」

    只断言「两个 xlsx 内容一样」不够：读回来是一样的表，字节可以不同（时间戳正是这样，
    它是元数据不是单元格）。要钉的恰恰是**字节**，所以比的是 git 自己算出来的哈希。
    """
    src = _isolate(monkeypatch, tmp_path)

    mf.build()
    first = _git(src, "rev-list", "master").stdout.split()
    mf.build()
    second = _git(src, "rev-list", "master").stdout.split()

    assert len(first) == 3, first
    assert first == second, (
        "同一个 build() 跑两次给出了不同的提交哈希 —— 造数不可复现，"
        "平台里已记录的那些身份（tip / 基线 / commits_log）重建一次就全废了。"
        "最先要查的是 xlsx 字节里随时间变化的东西（openpyxl 的时间戳、zip 条目时间）。"
    )
